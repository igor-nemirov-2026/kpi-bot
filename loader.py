"""Чтение шаблона KPI_import_template.xlsx и CSV-выгрузок из 1С."""
from __future__ import annotations

import csv
import datetime as dt
import shutil
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]


class TemplateError(Exception):
    """Ошибка в структуре шаблона — текст понятен пользователю."""


def _to_date(v) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return dt.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    raise TemplateError(f"Не понимаю дату «{s}». Используйте формат ДД.ММ.ГГГГ.")


def _to_num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if s in ("", "-", "—"):
        return None
    try:
        return float(s)
    except ValueError:
        raise TemplateError(f"Не число: «{v}»")


def _rows(ws, required: list[str]):
    """Строки листа как словари по заголовкам первой строки."""
    it = ws.iter_rows(values_only=True)
    try:
        head = [str(h).strip() if h is not None else "" for h in next(it)]
    except StopIteration:
        return
    missing = [c for c in required if c not in head]
    if missing:
        raise TemplateError(f"На листе «{ws.title}» нет столбцов: {', '.join(missing)}")
    for row in it:
        if row is None or all(v is None or v == "" for v in row):
            continue
        yield {h: row[i] if i < len(row) else None for i, h in enumerate(head) if h}


def _sheet(wb, name):
    if name not in wb.sheetnames:
        raise TemplateError(f"В файле нет листа «{name}». Не переименовывайте листы шаблона.")
    return wb[name]


def _month_table(wb, name) -> dict[str, dict[int, float]]:
    out: dict[str, dict[int, float]] = {}
    for r in _rows(_sheet(wb, name), ["kpi_id"] + MONTHS):
        kid = str(r["kpi_id"] or "").strip()
        if not kid:
            continue
        vals = {}
        for i, m in enumerate(MONTHS, 1):
            v = _to_num(r.get(m))
            if v is not None:
                vals[i] = v
        out[kid] = vals
    return out


def _read_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TemplateError(f"Не удалось прочитать кодировку файла {path.name}")
    dialect = ";" if text.count(";") >= text.count(",") else ","
    rows = list(csv.reader(text.splitlines(), delimiter=dialect))
    return [r for r in rows if any(c.strip() for c in r)]


def load(excel_path: Path, csv_dir: Path | None = None) -> dict:
    """Читает шаблон и CSV. Возвращает словарь с исходными данными."""
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise TemplateError(f"Не найден файл шаблона: {excel_path}")
    # копия — чтобы читать файл, даже если он открыт в Excel
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / "kpi.xlsx"
        shutil.copy2(excel_path, copy)
        wb = load_workbook(copy, data_only=True, read_only=True)

        settings = {}
        for r in _rows(_sheet(wb, "Настройки"), ["Параметр", "Значение"]):
            if r["Параметр"]:
                settings[str(r["Параметр"]).strip()] = r["Значение"]

        kpis = []
        for r in _rows(_sheet(wb, "Справочник_KPI"), ["kpi_id", "Роли", "Показатель", "Ед", "Направление", "Тип"]):
            kid = str(r["kpi_id"] or "").strip()
            if not kid:
                continue
            typ = str(r["Тип"] or "").strip().lower()
            if typ not in ("sum", "avg", "last", "formula"):
                raise TemplateError(f"Справочник_KPI, {kid}: тип «{typ}» — допустимо sum, avg, last, formula")
            kpis.append({
                "id": kid,
                "roles": [x.strip() for x in str(r.get("Роли") or "").split(",") if x.strip()],
                "section": str(r.get("Блок") or "").strip(),
                "name": str(r["Показатель"]).strip(),
                "unit": str(r.get("Ед") or "").strip(),
                "dir": "down" if str(r.get("Направление") or "").strip().lower() == "down" else "up",
                "type": typ,
                "formula": str(r.get("Формула") or "").strip(),
                "tile_roles": [x.strip() for x in str(r.get("Плитка") or "").split(",") if x.strip()],
                "best": str(r.get("Лучшая_практика") or "").strip(),
                "definition": str(r.get("Определение") or "").strip(),
                "source": str(r.get("Источник_1С") or "").strip(),
                "owner": str(r.get("Владелец") or "").strip(),
            })

        weights = {}
        for r in _rows(_sheet(wb, "Календарь"), ["Дата", "Рабочий"]):
            d = _to_date(r["Дата"])
            if not d:
                continue
            w = _to_num(r.get("Вес"))
            if w is None:  # формула без сохранённого значения
                w = _to_num(r.get("Рабочий")) or 0.0
            weights[d] = w

        plan = _month_table(wb, "План_месяцы")
        fact_m = _month_table(wb, "Факт_месяцы")
        last_year = _month_table(wb, "Прошлый_год")

        daily: dict[str, dict[dt.date, float]] = defaultdict(dict)
        for r in _rows(_sheet(wb, "Факт_дни"), ["Дата", "kpi_id", "Значение"]):
            d, kid, v = _to_date(r["Дата"]), str(r["kpi_id"] or "").strip(), _to_num(r["Значение"])
            if d and kid and v is not None:
                daily[kid][d] = v

        weekly: dict[str, dict[dt.date, float]] = defaultdict(dict)
        for r in _rows(_sheet(wb, "Факт_недели"), ["Неделя_с", "kpi_id", "Значение"]):
            d, kid, v = _to_date(r["Неделя_с"]), str(r["kpi_id"] or "").strip(), _to_num(r["Значение"])
            if d and kid and v is not None:
                weekly[kid][d - dt.timedelta(days=d.weekday())] = v
        wb.close()

    # CSV из 1С перекрывают значения листов (та же дата + kpi_id)
    csv_files = []
    if csv_dir and Path(csv_dir).exists():
        for path in sorted(Path(csv_dir).glob("fact_daily*.csv")):
            for row in _read_csv(path):
                if row[0].strip().lower() in ("дата", "date"):
                    continue
                d, kid, v = _to_date(row[0]), row[1].strip(), _to_num(row[2])
                if d and kid and v is not None:
                    daily[kid][d] = v
            csv_files.append(path.name)
        for path in sorted(Path(csv_dir).glob("fact_weekly*.csv")):
            for row in _read_csv(path):
                if row[0].strip().lower() in ("неделя_с", "week"):
                    continue
                d, kid, v = _to_date(row[0]), row[1].strip(), _to_num(row[2])
                if d and kid and v is not None:
                    weekly[kid][d - dt.timedelta(days=d.weekday())] = v
            csv_files.append(path.name)

    known = {k["id"] for k in kpis}
    unknown = sorted((set(daily) | set(weekly) | set(plan)) - known)
    return {
        "settings": settings, "kpis": kpis, "weights": weights, "plan": plan, "fact_m": fact_m,
        "last_year": last_year, "daily": dict(daily), "weekly": dict(weekly),
        "warnings": [f"kpi_id «{u}» нет в Справочник_KPI — строки пропущены" for u in unknown],
        "csv_files": csv_files,
    }
